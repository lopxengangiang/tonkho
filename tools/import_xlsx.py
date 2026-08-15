#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Nhập giá từ file Excel (bảng mã TCVN3) lên web tồn kho.

Cách dùng:
    python3 tools/import_xlsx.py Book1.xlsx                 # xem trước (dry-run)
    python3 tools/import_xlsx.py Book1.xlsx --execute       # ghi thật lên web

Mật khẩu lấy từ biến môi trường TONKHO_PASSWORD, hoặc nhập khi được hỏi.

Quy tắc merge (mã nào không có trong file thì giữ nguyên trên web):
- Cột dùng: Mã, Tên, ĐVT, Giá nhập → `nhap`, Giá lẻ → `ban`.
- Khớp mã chính xác trước; mã bị Excel cắt mất số 0 đầu (ô số) chỉ được khớp
  khi TÊN cũng khớp (tránh ghép nhầm sản phẩm khác trùng số).
- Giá 0/trống trong file = "không có giá" → giữ giá cũ, không ghi đè về 0.
- Mã mới có ít nhất một giá > 0 thì thêm mới; hàng không giá bị bỏ qua.
"""
import argparse, getpass, json, os, re, sys, unicodedata, urllib.request
from difflib import SequenceMatcher

import openpyxl

BASE = "https://tonkho.buicongminhhoang.workers.dev"
UA = "tonkho-import/1.0"

TCVN3 = {
    'µ':'à','¸':'á','¶':'ả','·':'ã','¹':'ạ',
    '¨':'ă','»':'ằ','¾':'ắ','¼':'ẳ','½':'ẵ','Æ':'ặ',
    '©':'â','Ç':'ầ','Ê':'ấ','È':'ẩ','É':'ẫ','Ë':'ậ',
    '®':'đ',
    'Ì':'è','Ð':'é','Î':'ẻ','Ï':'ẽ','Ñ':'ẹ',
    'ª':'ê','Ò':'ề','Õ':'ế','Ó':'ể','Ô':'ễ','Ö':'ệ',
    '×':'ì','Ý':'í','Ø':'ỉ','Ü':'ĩ','Þ':'ị',
    'ß':'ò','ã':'ó','á':'ỏ','â':'õ','ä':'ọ',
    '«':'ô','å':'ồ','è':'ố','æ':'ổ','ç':'ỗ','é':'ộ',
    '¬':'ơ','ê':'ờ','í':'ớ','ë':'ở','ì':'ỡ','î':'ợ',
    'ï':'ù','ó':'ú','ñ':'ủ','ò':'ũ','ô':'ụ',
    '\xad':'ư','õ':'ừ','ø':'ứ','ö':'ử','÷':'ữ','ù':'ự',
    'ú':'ỳ','ý':'ý','û':'ỷ','ü':'ỹ','þ':'ỵ',
    '¡':'Ă','¢':'Â','§':'Đ','£':'Ê','¤':'Ô','¥':'Ơ','¦':'Ư',
}

def conv(s):
    if s is None: return ""
    return "".join(TCVN3.get(c, c) for c in str(s)).strip()

def num(v):
    if v is None: return 0
    if isinstance(v, (int, float)): return int(round(v))
    s = re.sub(r'[^\d.-]', '', str(v))
    try: return int(round(float(s))) if s else 0
    except ValueError: return 0

def norm(s):
    s = unicodedata.normalize("NFD", str(s).lower())
    s = "".join(c for c in s if not unicodedata.combining(c)).replace("đ", "d")
    return re.sub(r'\s+', ' ', s).strip()

def ten_close(a, b):
    na, nb = norm(a), norm(b)
    if na == nb or na in nb or nb in na: return True
    return SequenceMatcher(None, na, nb).ratio() >= 0.75

def call(path, method="GET", body=None, token=None):
    req = urllib.request.Request(BASE + path, method=method,
        data=json.dumps(body).encode() if body else None,
        headers={"content-type": "application/json", "user-agent": UA,
                 **({"authorization": "Bearer " + token} if token else {})})
    with urllib.request.urlopen(req) as r:
        return json.load(r)

def read_excel(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = [conv(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {n: i for i, n in enumerate(head)}
    for want in ("Mã", "Tên", "Giá nhập", "Giá lẻ"):
        if want not in idx:
            sys.exit(f"Không thấy cột '{want}' — header đọc được: {head}")
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        ma = conv(r[idx["Mã"]])
        if not ma: continue
        rows[ma] = {  # trùng mã trong file: dòng cuối thắng
            "ma": ma, "ten": conv(r[idx["Tên"]]), "dvt": conv(r[idx.get("ĐVT", -1)]) if "ĐVT" in idx else "",
            "nhap": num(r[idx["Giá nhập"]]), "ban": num(r[idx["Giá lẻ"]]),
        }
    return rows

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--execute", action="store_true", help="ghi thật (mặc định chỉ xem trước)")
    args = ap.parse_args()

    password = os.environ.get("TONKHO_PASSWORD") or getpass.getpass("Mật khẩu web tồn kho: ")
    token = call("/api/login", "POST", {"password": password})["token"]
    live = call("/api/data", token=token)
    items = live["items"]
    by_ma = {it["ma"]: it for it in items}
    canon = {}
    for it in items:
        if it["ma"].isdigit():
            canon.setdefault(it["ma"].lstrip("0") or "0", []).append(it)

    excel = read_excel(args.xlsx)
    updated, added, skipped, ambiguous = [], [], [], []
    for ma, x in excel.items():
        target = by_ma.get(ma)
        if target is None and ma.isdigit():
            cands = [it for it in canon.get(ma.lstrip("0") or "0", []) if ten_close(x["ten"], it["ten"])]
            if len(cands) == 1: target = cands[0]
            elif len(cands) > 1: ambiguous.append((ma, [c["ma"] for c in cands])); continue
        if target is not None:
            ch = []
            for f in ("nhap", "ban"):
                if x[f] > 0 and x[f] != target[f]:
                    ch.append(f"{f}: {target[f]:,} → {x[f]:,}"); target[f] = x[f]
            if ch: updated.append((ma, target["ma"], target["ten"], ch))
        elif x["nhap"] > 0 or x["ban"] > 0:
            by_ma[ma] = dict(x); added.append(x)
        else:
            skipped.append(ma)

    print(f"File: {len(excel)} dòng | Web: {len(items)} mặt hàng")
    print(f"\nCẬP NHẬT GIÁ ({len(updated)}):")
    for ma, real, ten, ch in updated:
        via = "" if ma == real else f" (khớp {ma}→{real})"
        print(f"  {real:12} {ten[:42]:44}{via} {'; '.join(ch)}")
    print(f"\nTHÊM MỚI ({len(added)}):")
    for x in added:
        print(f"  {x['ma']:12} {x['ten'][:45]:47} nhập={x['nhap']:,} lẻ={x['ban']:,}")
    if skipped: print(f"\nBỎ QUA không giá: {skipped}")
    if ambiguous: sys.exit(f"\n⛔ Mã mập mờ, sửa file rồi chạy lại: {ambiguous}")

    if not args.execute:
        print("\n(dry-run — thêm --execute để ghi thật)"); return
    merged = list(by_ma.values())
    r = call("/api/data", "PUT", {"items": merged}, token=token)
    print(f"\nĐã ghi: {r}")

if __name__ == "__main__":
    main()
